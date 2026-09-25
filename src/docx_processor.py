from pathlib import Path
from docx import Document
import json


def load_document(file_path):
    """Load a DOCX document."""
    return Document(file_path)


def extract_paragraphs(document):
    """Extract non-empty paragraphs from the document."""
    paragraphs = []

    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()

        if text:
            paragraphs.append({
                "type": "paragraph",
                "index": index,
                "text": text
            })

    return paragraphs


def extract_tables(document):
    """Extract text from all tables in the document."""
    tables = []

    for table_index, table in enumerate(document.tables):
        table_data = []

        for row_index, row in enumerate(table.rows):
            row_data = []

            for cell_index, cell in enumerate(row.cells):
                text = cell.text.strip()

                row_data.append({
                    "row": row_index,
                    "column": cell_index,
                    "text": text
                })

            table_data.append(row_data)

        tables.append({
            "table_index": table_index,
            "rows": table_data
        })

    return tables


def extract_document_content(file_path):
    """Extract paragraphs and tables from a DOCX document."""
    document = load_document(file_path)

    return {
        "paragraphs": extract_paragraphs(document),
        "tables": extract_tables(document)
    }


def save_extracted_content(content, output_file):
    """Save extracted document content as JSON."""
    with open(output_file, "w", encoding="utf-8") as file:
        json.dump(content, file, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    input_file = Path("input/Red Herring Prospectus.docx")
    output_file = Path("extracted/document_content.json")

    if not input_file.exists():
        raise FileNotFoundError(
            f"Input file not found: {input_file}"
        )

    content = extract_document_content(input_file)

    save_extracted_content(content, output_file)

    print("Document loaded successfully.")
    print(f"Paragraphs found: {len(content['paragraphs'])}")
    print(f"Tables found: {len(content['tables'])}")
    print(f"Extracted content saved to: {output_file}")